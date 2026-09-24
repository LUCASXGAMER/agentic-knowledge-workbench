from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path


@dataclass
class ParsedPage:
    text: str
    page: int | None = None
    section_title: str = ""


@dataclass
class ParsedDocument:
    pages: list[ParsedPage]
    message: str = "解析成功"


def parse_document(path: str | Path) -> ParsedDocument:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        return ParsedDocument([ParsedPage(file_path.read_text(encoding="utf-8"), 1)])
    if suffix == ".csv":
        return _parse_csv(file_path)
    if suffix == ".docx":
        return _parse_docx(file_path)
    if suffix == ".xlsx":
        return _parse_xlsx(file_path)
    if suffix == ".pdf":
        return _parse_pdf(file_path)
    if suffix in {".png", ".jpg", ".jpeg"}:
        return ParsedDocument([ParsedPage("图片内容待文字识别，请开启 OCR 后重新入库。", 1)], "待文字识别")
    return ParsedDocument([ParsedPage(file_path.read_text(encoding="utf-8", errors="ignore"), 1)])


def _parse_csv(file_path: Path) -> ParsedDocument:
    rows: list[str] = []
    with file_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for index, row in enumerate(reader, start=1):
            values = "；".join(f"{key}：{value}" for key, value in row.items())
            rows.append(f"第 {index} 行：{values}")
    return ParsedDocument([ParsedPage("\n".join(rows), 1, file_path.stem)])


def _parse_docx(file_path: Path) -> ParsedDocument:
    try:
        from docx import Document as DocxDocument
    except Exception:
        return ParsedDocument([ParsedPage("Word 解析能力未启用，请安装对应解析组件后重新入库。", 1)], "Word 解析能力未启用")
    doc = DocxDocument(str(file_path))
    lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return ParsedDocument([ParsedPage("\n".join(lines), 1, file_path.stem)])


def _parse_xlsx(file_path: Path) -> ParsedDocument:
    try:
        from openpyxl import load_workbook
    except Exception:
        return ParsedDocument([ParsedPage("Excel 解析能力未启用，请安装对应解析组件后重新入库。", 1)], "Excel 解析能力未启用")
    workbook = load_workbook(file_path, data_only=True, read_only=True)
    pages: list[ParsedPage] = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(cell or "").strip() for cell in rows[0]]
        lines: list[str] = []
        for row_index, row in enumerate(rows[1:], start=2):
            cells = []
            for col_index, value in enumerate(row):
                header = headers[col_index] if col_index < len(headers) and headers[col_index] else f"列{col_index + 1}"
                cells.append(f"{header}：{value}")
            lines.append(f"{sheet.title} 第 {row_index} 行，" + "；".join(cells))
        pages.append(ParsedPage("\n".join(lines), None, sheet.title))
    return ParsedDocument(pages or [ParsedPage("", 1, file_path.stem)])


def _parse_pdf(file_path: Path) -> ParsedDocument:
    try:
        import fitz
    except Exception:
        return ParsedDocument([ParsedPage("PDF 解析能力未启用，请安装对应解析组件后重新入库。", 1)], "PDF 解析能力未启用")
    pages: list[ParsedPage] = []
    with fitz.open(file_path) as doc:
        for index, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if len(text) < 30:
                text = text or "本页文字较少，请开启 OCR 后重新入库。"
            pages.append(ParsedPage(text, index, file_path.stem))
    return ParsedDocument(pages, "解析成功")
